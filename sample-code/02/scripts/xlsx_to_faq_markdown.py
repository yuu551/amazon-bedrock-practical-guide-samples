"""FAQ表(faq.xlsx)からKnowledge Bases取り込み用Markdownを一括生成する。

使い方:
    uv run python scripts/xlsx_to_faq_markdown.py
    uv run python scripts/xlsx_to_faq_markdown.py --input <path>.xlsx --output <dir>

入力xlsxは1シート目に次の列を想定する。
    タイトル, 質問, 回答, 補足, カテゴリ, 対象者

「タイトル」「質問」「回答」を必須、「補足」は任意。
「カテゴリ」「対象者」は表上の管理項目で、Markdown本文には載せない
（メタデータ付与は第3章でAIに任せる流れに合わせるため）。
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

DEFAULT_INPUT = Path(__file__).resolve().parents[1] / "sample-data" / "faq" / "source" / "faq.xlsx"
DEFAULT_OUTPUT = Path(__file__).resolve().parents[1] / "sample-data" / "faq"

REQUIRED_COLUMNS = ("タイトル", "質問", "回答")
OPTIONAL_COLUMNS = ("補足", "カテゴリ", "対象者")


def read_rows(xlsx_path: Path) -> list[dict[str, str]]:
    if not xlsx_path.exists():
        raise FileNotFoundError(f"入力ファイルが見つかりません: {xlsx_path}")
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet = workbook.active

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return []
    headers = [str(cell).strip() if cell is not None else "" for cell in header]
    missing = [col for col in REQUIRED_COLUMNS if col not in headers]
    if missing:
        raise ValueError(f"必須列が見つかりません: {missing} (実際の列: {headers})")

    rows: list[dict[str, str]] = []
    for raw in rows_iter:
        if raw is None or all(cell is None or str(cell).strip() == "" for cell in raw):
            continue
        record = {headers[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(raw) if i < len(headers)}
        for col in OPTIONAL_COLUMNS:
            record.setdefault(col, "")
        rows.append(record)
    return rows


def render_markdown(row: dict[str, str]) -> str:
    title = row["タイトル"]
    question = row["質問"]
    answer = row["回答"]
    supplementary = row.get("補足", "")

    lines = [f"# {title}", "", f"質問: {question}", "", f"回答: {answer}"]
    if supplementary:
        lines.extend(["", f"補足: {supplementary}"])
    return "\n".join(lines) + "\n"


def filename_for(title: str) -> str:
    return title.replace(" ", "_") + ".md"


def write_markdown(rows: list[dict[str, str]], output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for row in rows:
        if not row.get("タイトル"):
            raise ValueError(f"タイトルが空の行があります: {row}")
        path = output_dir / filename_for(row["タイトル"])
        path.write_text(render_markdown(row), encoding="utf-8")
        written.append(path)
    return written


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="FAQ表をMarkdownに変換する")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="入力xlsxのパス")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="出力ディレクトリ")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    rows = read_rows(args.input)
    written = write_markdown(rows, args.output)
    print(f"{len(written)} 件のMarkdownを {args.output} に出力しました")
    return 0


if __name__ == "__main__":
    sys.exit(main())
