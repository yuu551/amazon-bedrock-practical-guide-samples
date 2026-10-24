from __future__ import annotations

import json
from pathlib import Path
from statistics import mean
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Excel list-cell conventions
#   - List values are stored as semicolon-separated strings in a single cell
#   - Example: "item1; item2; item3"
# ---------------------------------------------------------------------------

_LIST_SEP = ";"


def split_list_cell(value: Any) -> list[str]:
    if value is None or str(value).strip() == "":
        return []
    return [item.strip() for item in str(value).split(_LIST_SEP) if item.strip()]


def join_list_cell(items: list[str]) -> str:
    return "; ".join(items)


# ---------------------------------------------------------------------------
# Excel styling helpers
# ---------------------------------------------------------------------------

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill(
    start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
)
_SCORE_FMT = "0.0"


def style_header_row(ws: Any, headers: list[str]) -> None:
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def auto_fit_columns(ws: Any, min_width: int = 10, max_width: int = 50) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        lengths: list[int] = []
        for cell in col:
            if cell.value:
                lines = str(cell.value).split("\n")
                lengths.append(max(len(line) for line in lines))
        width = max(lengths) if lengths else min_width
        ws.column_dimensions[letter].width = min(max(width + 2, min_width), max_width)


def apply_score_color_scale(
    ws: Any,
    score_columns: tuple[str, ...] = ("E", "F", "G"),
    start_row: int = 2,
    end_row: int | None = None,
) -> None:
    """正確性・完全性・根拠性の3列に3色スケール（0=赤 / 0.5=黄 / 1.0=緑）を適用する。

    eval_workbook.xlsx の評価結果シートでは E〜G 列が 正確性/完全性/根拠性 のスコアに当たる。
    """
    from openpyxl.formatting.rule import ColorScaleRule
    last = end_row if end_row is not None else ws.max_row
    if last < start_row:
        return
    for col in score_columns:
        rule = ColorScaleRule(
            start_type="num", start_value=0.0, start_color="F8696B",
            mid_type="num", mid_value=0.5, mid_color="FFEB84",
            end_type="num", end_value=1.0, end_color="63BE7B",
        )
        ws.conditional_formatting.add(f"{col}{start_row}:{col}{last}", rule)


def _open_or_create_workbook(path: Path) -> Workbook:
    if path.exists():
        return load_workbook(path)
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    return wb


def _read_sheet_as_dicts(
    path: Path, sheet_name: str
) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if len(rows) < 2:
        return []
    headers = [str(h).strip() for h in rows[0] if h is not None]
    entries: list[dict[str, Any]] = []
    for row in rows[1:]:
        if not any(cell is not None for cell in row):
            continue
        entries.append(dict(zip(headers, row)))
    return entries


# ---------------------------------------------------------------------------
# 評価ワークブック I/O  (シート: 評価ケース, 採点基準, 評価結果)
# ---------------------------------------------------------------------------

def read_eval_cases(
    path: Path, sheet_name: str = "評価ケース"
) -> list[dict[str, Any]]:
    cases: list[dict[str, Any]] = []
    for raw in _read_sheet_as_dicts(path, sheet_name):
        case: dict[str, Any] = {
            "id": str(raw.get("ケースID", "")),
            "question": str(raw.get("質問", "")),
            "expected_points": split_list_cell(raw.get("期待要点")),
            "expected_sources": split_list_cell(raw.get("期待ソース")),
        }
        filter_val = raw.get("フィルター条件")
        if filter_val and str(filter_val).strip():
            try:
                case["filter"] = json.loads(str(filter_val))
            except json.JSONDecodeError:
                pass
        cases.append(case)
    return cases


def read_rubrics(
    path: Path, sheet_name: str = "採点基準"
) -> dict[str, Any]:
    rubrics: dict[str, Any] = {}
    for raw in _read_sheet_as_dicts(path, sheet_name):
        metric = str(raw.get("評価項目", "")).strip()
        if not metric:
            continue
        rubrics[metric] = {
            "満点": str(raw.get("満点", "")),
            "部分点": str(raw.get("部分点", "")),
            "不合格": str(raw.get("不合格", "")),
        }
    return rubrics


_RESULTS_HEADERS = [
    "ケースID",
    "質問",
    "回答",
    "参照先",
    "正確性",
    "完全性",
    "根拠性",
    "主な失敗類型",
    "採点理由",
    "ソースヒット率",
    "応答時間(ms)",
    "エラー",
]


def write_results_sheet(
    path: Path,
    results: list[dict[str, Any]],
    sheet_name: str = "評価結果",
) -> None:
    wb = _open_or_create_workbook(path)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    style_header_row(ws, _RESULTS_HEADERS)

    for r in results:
        scores = r.get("scores", {})
        reasons = r.get("reason_by_metric", {})
        reason_parts = [f"{k}: {v}" for k, v in reasons.items() if v]
        reason_text = (
            "\n".join(reason_parts) if reason_parts else r.get("overall_summary", "")
        )
        ws.append(
            [
                r.get("id", ""),
                r.get("question", ""),
                r.get("answer", ""),
                join_list_cell(r.get("references", [])),
                scores.get("correctness", 0.0),
                scores.get("completeness", 0.0),
                scores.get("groundedness", 0.0),
                r.get("primary_failure_type", ""),
                reason_text,
                r.get("expected_source_hit_ratio"),
                r.get("latency_ms", 0),
                r.get("error", ""),
            ]
        )

    for row in ws.iter_rows(min_row=2, min_col=5, max_col=7):
        for cell in row:
            cell.number_format = _SCORE_FMT

    auto_fit_columns(ws)
    # 正確性・完全性・根拠性に3色スケール条件付き書式を再付与
    apply_score_color_scale(ws)
    wb.save(path)


# ---------------------------------------------------------------------------
# フィードバックワークブック I/O  (シート: フィードバック, 更新候補)
# ---------------------------------------------------------------------------

def read_feedback_log(
    path: Path, sheet_name: str = "フィードバック"
) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    for raw in _read_sheet_as_dicts(path, sheet_name):
        entries.append(
            {
                "feedback_id": str(raw.get("フィードバックID", "")),
                "question": str(raw.get("質問", "")),
                "agent_answer": str(raw.get("自動回答", "")),
                "operator_fix": str(raw.get("オペレーター修正", "")),
                "comment": str(raw.get("補足コメント", "")),
                "resolved_by_human": str(raw.get("有人対応", "")).upper()
                in ("TRUE", "1", "YES"),
                "related_case_ids": split_list_cell(raw.get("関連ケースID")),
            }
        )
    return entries


_CANDIDATES_HEADERS = [
    "フィードバックID",
    "アクション",
    "対象文書",
    "欠落点",
    "変更案",
    "理由",
    "再評価ケースID",
]


def write_candidates_sheet(
    path: Path,
    candidates: list[dict[str, Any]],
    sheet_name: str = "更新候補",
) -> None:
    wb = _open_or_create_workbook(path)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    style_header_row(ws, _CANDIDATES_HEADERS)

    for c in candidates:
        ws.append(
            [
                c.get("feedback_id", ""),
                c.get("action", ""),
                c.get("target_document", ""),
                join_list_cell(c.get("missing_points", [])),
                c.get("draft_change", ""),
                c.get("reason", ""),
                join_list_cell(c.get("retest_case_ids", [])),
            ]
        )

    auto_fit_columns(ws)
    wb.save(path)


# ---------------------------------------------------------------------------
# 承認済み更新ワークブック I/O  (シート: 承認済み更新)
# ---------------------------------------------------------------------------

def read_approved_updates(
    path: Path, sheet_name: str = "承認済み更新"
) -> list[dict[str, Any]]:
    updates: list[dict[str, Any]] = []
    for raw in _read_sheet_as_dicts(path, sheet_name):
        update: dict[str, Any] = {
            "key": str(raw.get("S3キー", "")),
            "body": str(raw.get("本文", "")),
        }
        metadata_val = raw.get("メタデータJSON")
        if metadata_val and str(metadata_val).strip():
            try:
                update["metadata_payload"] = json.loads(str(metadata_val))
            except json.JSONDecodeError:
                pass
        updates.append(update)
    return updates


# ---------------------------------------------------------------------------
# Bedrock Converse / RetrieveAndGenerate response helpers
# ---------------------------------------------------------------------------

# 温度を受け付けず、適応的思考が既定で有効なモデル。
# モデルIDと推論プロファイルARNのどちらでも末尾で判定する。
REASONING_MODELS = (
    "anthropic.claude-sonnet-5",
    "anthropic.claude-sonnet-5-5",
    "anthropic.claude-opus-5",
    "anthropic.claude-opus-5-5",
    "anthropic.claude-haiku-5-5",
    "openai.gpt-6-sol",
    "openai.gpt-6.1-sol",
    "openai.gpt-6-luna",
)


def uses_reasoning_defaults(model_id: str | None) -> bool:
    return (model_id or "").endswith(REASONING_MODELS)


def converse_inference_config(
    model_id: str | None, max_tokens: int, temperature: float
) -> dict[str, Any]:
    """新しいモデルには温度を送らず、思考で使う分も出力上限に含める。"""
    if uses_reasoning_defaults(model_id):
        return {"maxTokens": max(max_tokens, 4096)}
    return {"maxTokens": max_tokens, "temperature": temperature}


# 新しいモデルではKB既定のプロンプトが用意されていない場合がある。
# 引用を返すため、AWSが展開する出力形式の指示も保持する。
KB_GENERATION_PROMPT = (
    "検索結果に基づいて日本語で回答してください。資料にないことは推測せず、不明と伝えてください。"
    "資料の適用時期が異なる場合は区別してください。\n"
    "質問: $query$\n検索結果:\n$search_results$\n$output_format_instructions$"
)
KB_ORCHESTRATION_PROMPT = (
    "会話履歴と質問から、関連資料を探す検索クエリを作成してください。\n"
    "会話履歴:\n$conversation_history$\n質問: $query$\n$output_format_instructions$"
)


def knowledge_base_prompts(model_id: str | None) -> dict[str, Any]:
    """新しいモデルのときだけ、RetrieveAndGenerateに渡す2つのプロンプトを返す。"""
    if not uses_reasoning_defaults(model_id):
        return {}
    return {
        "generationConfiguration": {
            "promptTemplate": {"textPromptTemplate": KB_GENERATION_PROMPT},
            "inferenceConfig": {"textInferenceConfig": {"maxTokens": 4096}},
        },
        "orchestrationConfiguration": {
            "promptTemplate": {"textPromptTemplate": KB_ORCHESTRATION_PROMPT},
            "inferenceConfig": {"textInferenceConfig": {"maxTokens": 4096}},
        },
    }


def extract_text_from_converse_response(response: dict[str, Any]) -> str:
    content = response.get("output", {}).get("message", {}).get("content", [])
    parts: list[str] = []
    for block in content:
        text = block.get("text")
        if text:
            parts.append(text)
    return "\n".join(parts).strip()


def extract_json_object(text: str) -> dict[str, Any]:
    start = text.find("{")
    if start == -1:
        raise ValueError("No JSON object found in model output.")

    depth = 0
    in_string = False
    escape = False
    end = None
    for index in range(start, len(text)):
        char = text[index]
        if escape:
            escape = False
            continue
        if char == "\\":
            escape = True
            continue
        if char == '"':
            in_string = not in_string
            continue
        if in_string:
            continue
        if char == "{":
            depth += 1
        elif char == "}":
            depth -= 1
            if depth == 0:
                end = index + 1
                break

    if end is None:
        raise ValueError("JSON object appears to be incomplete.")

    return json.loads(text[start:end])


def extract_retrieved_references(
    response: dict[str, Any],
) -> list[dict[str, Any]]:
    references: list[dict[str, Any]] = []
    for citation in response.get("citations", []):
        references.extend(citation.get("retrievedReferences", []))
    return references


_LOCATION_KEYS = [
    ("s3Location", "uri"),
    ("webLocation", "url"),
    ("confluenceLocation", "url"),
    ("salesforceLocation", "url"),
    ("sharePointLocation", "url"),
    ("kendraDocumentLocation", "uri"),
    ("customDocumentLocation", "id"),
    ("sqlLocation", "query"),
]


def reference_location(reference: dict[str, Any]) -> str:
    location = reference.get("location", {})
    for loc_type, key in _LOCATION_KEYS:
        if loc_type in location:
            return location[loc_type].get(key, "")
    return ""


def expected_source_hit_ratio(
    references: list[str],
    expected_sources: list[str] | None,
) -> float | None:
    if not expected_sources:
        return None
    lowered_refs = [r.lower() for r in references]
    hits = sum(
        1
        for es in expected_sources
        if any(es.lower() in r for r in lowered_refs)
    )
    return hits / len(expected_sources)


# ---------------------------------------------------------------------------
# Markdown summary helper
# ---------------------------------------------------------------------------

def markdown_table(headers: list[str], rows: list[list[str]]) -> str:
    table = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        table.append("| " + " | ".join(row) + " |")
    return "\n".join(table)


def safe_mean(values: list[float]) -> float | None:
    if not values:
        return None
    return float(mean(values))
