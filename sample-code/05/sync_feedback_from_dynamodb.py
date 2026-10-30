#!/usr/bin/env python3
"""Sync user feedback from DynamoDB to feedback_workbook.xlsx.

第3章のRAGポータルがAmplify Data経由でDynamoDBへ書き込んだフィードバックを
取得し、第5章のフィードバックワークブック（``data/feedback_workbook.xlsx``）の
``フィードバック``シートに変換して記録します。出力はそのまま
``extract_feedback_candidates.py`` の入力として使えます。

DynamoDBテーブル名は ``amplify_outputs.json`` の ``data.tables`` または
AWSコンソールから確認できます。実行時は ``--table-name`` または
``FEEDBACK_TABLE_NAME`` 環境変数で指定してください。
"""
from __future__ import annotations

import argparse
import os
from pathlib import Path
from typing import Any, Iterator

import boto3

from common import (
    auto_fit_columns,
    join_list_cell,
    style_header_row,
)
from openpyxl import Workbook


_FEEDBACK_HEADERS = [
    "フィードバックID",
    "質問",
    "自動回答",
    "オペレーター修正",
    "補足コメント",
    "有人対応",
    "関連ケースID",
    "評価",
    "検索プリセット",
    "部門フィルタ",
    "応答時間(ms)",
    "参照URI",
    "受信日時",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Sync feedback from DynamoDB to feedback_workbook.xlsx"
    )
    parser.add_argument(
        "--table-name",
        default=os.environ.get("FEEDBACK_TABLE_NAME"),
        help="Amplify DataがプロビジョニングしたDynamoDBテーブル名（Feedback-xxxx-NONE）。",
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=Path("data/feedback_workbook.xlsx"),
        help="記録先のExcelワークブック。",
    )
    parser.add_argument(
        "--region",
        default=os.environ.get("AWS_REGION", "us-east-1"),
    )
    return parser.parse_args()


def scan_all(table: Any) -> Iterator[dict[str, Any]]:
    response = table.scan()
    yield from response.get("Items", [])
    while "LastEvaluatedKey" in response:
        response = table.scan(ExclusiveStartKey=response["LastEvaluatedKey"])
        yield from response.get("Items", [])


def to_feedback_row(item: dict[str, Any]) -> list[Any]:
    rating = item.get("rating")
    citations = item.get("citations") or []

    return [
        item.get("id", ""),                     # フィードバックID（DynamoDBのidをそのまま使う）
        item.get("question", ""),               # 質問
        item.get("agentAnswer", ""),            # 自動回答
        "",                                     # オペレーター修正（運用フローで埋める）
        item.get("comment", "") or "",          # 補足コメント
        "FALSE",                                # 有人対応（運用フローで埋める）
        "",                                     # 関連ケースID（あとで紐づけ）
        rating or "",                           # 評価（UP / DOWN）
        item.get("searchPresetLabel", "") or "",
        item.get("filterDepartment", "") or "",
        item.get("latencyMs", "") or "",
        join_list_cell([str(c) for c in citations]),
        item.get("createdAt", "") or "",
    ]


def write_feedback_workbook(path: Path, rows: list[list[Any]]) -> None:
    """既存のExcelに依存せず、Amplifyから来たデータで上書きする。

    ``extract_feedback_candidates.py`` は ``フィードバック`` シートを `read_feedback_log` で
    読み出すだけなので、ここで全件作り直してしまって問題ない。
    オペレーター修正や関連ケースIDを手動で書き加える運用なら、運用後に
    Excelで上書きしたものを反映する形も検討する。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "フィードバック"
    style_header_row(ws, _FEEDBACK_HEADERS)
    for row in rows:
        ws.append(row)
    auto_fit_columns(ws)

    # 更新候補シートは extract_feedback_candidates.py が再生成するため、空シートとして残す
    ws_c = wb.create_sheet("更新候補")
    style_header_row(
        ws_c,
        ["フィードバックID", "アクション", "対象文書", "欠落点",
         "変更案", "理由", "再評価ケースID"],
    )
    auto_fit_columns(ws_c)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    args = parse_args()
    if not args.table_name:
        raise SystemExit(
            "Provide --table-name or FEEDBACK_TABLE_NAME (Amplify DataのDynamoDBテーブル名)."
        )

    dynamodb = boto3.resource("dynamodb", region_name=args.region)
    table = dynamodb.Table(args.table_name)

    rows: list[list[Any]] = []
    for item in scan_all(table):
        rows.append(to_feedback_row(item))

    write_feedback_workbook(args.workbook, rows)
    print(f"Wrote {len(rows)} feedback entries to '{args.workbook}'")


if __name__ == "__main__":
    main()
