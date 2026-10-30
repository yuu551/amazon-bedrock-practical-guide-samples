#!/usr/bin/env python3
"""Generate sample Excel workbook for Chapter 5 agent evaluation.

Creates ``agent_data/agent_eval_workbook.xlsx`` with three sheets:

- ``評価ケース``  — case ID, question, expected points, expected trajectory,
  forbidden tools, category.
- ``採点基準``  — custom rubrics for OutputAccuracy / Trajectory and guidance
  for interpreting ToolSelection / GoalSuccess.
- ``評価結果``  — empty template that ``run_agent_batch_eval.py`` fills.

Run once before executing the agent evaluation script. Readers can then edit
the workbook in Excel before re-running.
"""
from __future__ import annotations

from pathlib import Path

from common import auto_fit_columns, style_header_row

from openpyxl import Workbook

DATA_DIR = Path(__file__).parent / "agent_data"


def create_agent_eval_workbook() -> None:
    wb = Workbook()

    # --- 評価ケース ---
    ws = wb.active
    ws.title = "評価ケース"
    style_header_row(
        ws,
        [
            "ケースID",
            "質問",
            "期待要点",
            "期待軌跡",
            "不要ツール",
            "カテゴリ",
        ],
    )
    cases = [
        # ケースA: 鮮度の高い情報源を取り損ねるケース（ツール選択ミス）
        [
            "agent-incident-502",
            "APIで`HTTP 502 Bad Gateway`が断続的に発生しています。"
            "プラットフォーム側の既知障害ですか。",
            "HTTP 502が既知障害であると案内する; "
            "5MBを超えるPOSTリクエストで発生しやすいと案内する; "
            "回避策として5MB以下への分割、5秒後の再試行（最大3回）、gzip圧縮を案内する; "
            "ロードバランサーのタイムアウト設定を調整予定と案内する",
            "search_support_kb; confluence_search",
            "",
            "tool-selection-miss",
        ],
        # ケースB: silent failure（回答は正しいが過剰なツール呼び出し）
        [
            "agent-sla-credit-noise",
            "Enterpriseプランで稼働率が99.5%まで落ちている日が続いています。"
            "サービスクレジットの計算方法を知りたいので、"
            "サポートに連絡したいです。",
            "クレジット率（99.0%以上99.95%未満は10%、"
            "95.0%以上99.0%未満は25%、95.0%未満は50%）を案内する; "
            "申請窓口として管理コンソールまたはサポート窓口のいずれかを案内する; "
            "申請期限（障害発生月の翌月末）に触れる",
            "search_support_kb",
            "confluence_search; create_jira_ticket",
            "silent-failure-overcall",
        ],
        # 補強: 単純FAQで過剰起票してしまうケース
        [
            "agent-saml-login-noise",
            "SAML SSOでログインできない顧客がいます。"
            "最初に確認すべき項目を教えてください。",
            "IdP側のユーザー割当を確認する; "
            "ACS URLとEntity ID（SP EntityID）の一致を確認する; "
            "IdP証明書の有効期限を確認する; "
            "SAMLログ（直近30件）の確認を案内する",
            "search_support_kb",
            "confluence_search; create_jira_ticket",
            "silent-failure-overcall",
        ],
    ]
    for row in cases:
        ws.append(row)
    auto_fit_columns(ws)

    # --- 採点基準 ---
    ws_r = wb.create_sheet("採点基準")
    style_header_row(ws_r, ["評価項目", "満点", "部分点", "不合格"])
    rubrics = [
        [
            "OutputAccuracy",
            "期待要点をすべて満たし、事実誤認がない",
            "一部の要点が欠落している、または細部に誤りがある",
            "主要な要点が欠落している、または重要な事実誤認がある",
        ],
        [
            "ToolSelection",
            "実際に呼んだ各ツールが、その時点の質問と文脈に対して妥当",
            "一部のツール呼び出しだけが妥当",
            "実際に呼んだツールが質問や文脈に対して妥当でない",
        ],
        [
            "Trajectory",
            "期待軌跡と完全一致（順序・回数とも）",
            "期待ツールは網羅するが順序が違う、"
            "または余計な呼び出しが混じる",
            "期待ツールが抜けている、または明らかに無駄な経路を辿る",
        ],
        [
            "GoalSuccess",
            "ユーザーゴールを達成し、副作用（不要起票・誤誘導）がない",
            "ゴールは達成したが不要起票や軽微な誤誘導がある",
            "ゴール未達成",
        ],
    ]
    for row in rubrics:
        ws_r.append(row)
    auto_fit_columns(ws_r, max_width=60)

    # --- 評価結果 (空テンプレート、run_agent_batch_eval.py が記録する) ---
    ws_res = wb.create_sheet("評価結果")
    style_header_row(
        ws_res,
        [
            "ケースID",
            "質問",
            "回答",
            "実際の軌跡",
            "不要呼び出し数",
            "OutputAccuracy",
            "ToolSelection",
            "Trajectory",
            "GoalSuccess",
            "主な失敗類型",
            "採点理由",
            "応答時間(ms)",
            "エラー",
        ],
    )
    auto_fit_columns(ws_res)

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    path = DATA_DIR / "agent_eval_workbook.xlsx"
    wb.save(path)
    print(f"Created {path}  ({len(cases)} cases)")


if __name__ == "__main__":
    create_agent_eval_workbook()
