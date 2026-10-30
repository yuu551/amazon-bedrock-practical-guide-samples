#!/usr/bin/env python3
"""Re-judge existing batch eval answers without re-calling the application API.

Reads stored answer/references from the '評価結果' sheet of `data/eval_workbook.xlsx`
and re-scores each case with the specified LLM judge. Writes scores back to the
same workbook and refreshes `output/summary.md`.

Use when the answer-generation pipeline has not changed but the judge model
needs to be swapped (e.g. Sonnet 4.5 → Sonnet 4.6 unification).
"""
from __future__ import annotations

import argparse
import os
import time
from pathlib import Path
from typing import Any

import boto3
from openpyxl import load_workbook

from common import (
    expected_source_hit_ratio,
    read_eval_cases,
    read_rubrics,
    split_list_cell,
    write_results_sheet,
)
from run_batch_eval import judge_case, summarize


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--workbook",
        type=Path,
        default=Path("data/eval_workbook.xlsx"),
    )
    parser.add_argument(
        "--source-sheet",
        default="評価結果",
        help="Sheet to read existing answers/references from.",
    )
    parser.add_argument(
        "--judge-model-id",
        default=os.environ.get("JUDGE_MODEL_ID", "us.anthropic.claude-sonnet-4-6"),
    )
    parser.add_argument(
        "--region",
        default=os.environ.get("AWS_REGION", "us-east-1"),
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("output"),
    )
    return parser.parse_args()


def read_existing_results(path: Path, sheet_name: str) -> dict[str, dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if len(rows) < 2:
        return {}
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out: dict[str, dict[str, Any]] = {}
    for row in rows[1:]:
        if not any(c is not None for c in row):
            continue
        record = dict(zip(headers, row))
        case_id = str(record.get("ケースID", "")).strip()
        if not case_id:
            continue
        out[case_id] = {
            "answer": record.get("回答", "") or "",
            "references": split_list_cell(record.get("参照先")),
            "latency_ms": float(record.get("応答時間(ms)") or 0),
        }
    return out


def main() -> None:
    args = parse_args()

    cases = read_eval_cases(args.workbook)
    rubrics = read_rubrics(args.workbook)
    existing = read_existing_results(args.workbook, args.source_sheet)
    runtime = boto3.client("bedrock-runtime", region_name=args.region)

    results: list[dict[str, Any]] = []
    for case in cases:
        prior = existing.get(case["id"])
        if prior is None:
            print(f"[skip] {case['id']}: 既存の回答が見つかりません")
            continue

        answer = str(prior["answer"])
        references = prior["references"]
        started = time.perf_counter()
        try:
            judge_result = judge_case(
                runtime,
                judge_model_id=args.judge_model_id,
                case=case,
                answer=answer,
                references=references,
                rubrics=rubrics,
            )
            error = None
        except Exception as exc:  # noqa: BLE001
            judge_result = {
                "scores": {"correctness": 0.0, "completeness": 0.0, "groundedness": 0.0},
                "reason_by_metric": {},
                "overall_summary": "",
                "primary_failure_type": "none",
            }
            error = str(exc)
        judge_elapsed_ms = round((time.perf_counter() - started) * 1000, 1)

        results.append(
            {
                "id": case["id"],
                "question": case["question"],
                "answer": answer,
                "references": references,
                "scores": judge_result["scores"],
                "reason_by_metric": judge_result.get("reason_by_metric", {}),
                "overall_summary": judge_result.get("overall_summary", ""),
                "primary_failure_type": judge_result.get("primary_failure_type", "none"),
                "expected_source_hit_ratio": expected_source_hit_ratio(
                    references, case.get("expected_sources")
                ),
                "latency_ms": prior["latency_ms"],
                "error": error,
            }
        )
        s = judge_result["scores"]
        print(
            f"[{case['id']}] correctness={s['correctness']:.1f} "
            f"completeness={s['completeness']:.1f} groundedness={s['groundedness']:.1f} "
            f"judge={judge_elapsed_ms:.0f}ms"
        )

    write_results_sheet(args.workbook, results)
    print(f"Wrote {len(results)} results to '{args.workbook}' (評価結果シート)")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    summary_path = args.output_dir / "summary.md"
    summary_path.write_text(summarize(results), encoding="utf-8")
    print(f"Wrote summary to {summary_path}")


if __name__ == "__main__":
    main()
